"""Offline resume text extraction for TXT, DOCX and PDF files."""

from __future__ import annotations

from pathlib import Path
import re
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from pydantic import BaseModel, Field


class ResumeParseError(ValueError):
    """Raised when a resume cannot be parsed safely."""


class ParsedResume(BaseModel):
    """Format-independent result consumed by later workflow stages."""

    source_name: str
    file_type: str
    text: str = Field(min_length=1)
    unit_count: int = Field(ge=1)
    character_count: int = Field(ge=1)


SUPPORTED_SUFFIXES = {".txt", ".docx", ".pdf"}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 50_000
MAX_PDF_PAGES = 50
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def _clean_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _read_txt(path: Path) -> tuple[str, int]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            cleaned = _clean_text(text)
            return cleaned, max(1, len(cleaned.splitlines()))
        except UnicodeDecodeError:
            continue
    raise ResumeParseError("无法识别 TXT 编码；请另存为 UTF-8 后重试。")


def _read_docx(path: Path) -> tuple[str, int]:
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ResumeParseError("文件不是有效的 DOCX 文档。")
            total_size = sum(item.file_size for item in archive.infolist())
            if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeParseError("DOCX 解压后内容过大，已拒绝解析。")
    except BadZipFile as exc:
        raise ResumeParseError("文件不是有效的 DOCX 文档。") from exc

    try:
        from docx import Document
        document = Document(path)
    except ImportError as exc:
        raise ResumeParseError("缺少 python-docx，请先安装项目依赖。") from exc
    except Exception as exc:
        raise ResumeParseError(f"DOCX 文件损坏或无法读取：{exc}") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    cleaned_blocks = [_clean_text(block) for block in blocks]
    nonempty = [block for block in cleaned_blocks if block]

    # Word text boxes and drawing objects are stored in XML nodes that
    # python-docx does not expose through document.paragraphs. Use a
    # read-only OOXML fallback only when the public API found no text.
    if not nonempty:
        try:
            with ZipFile(path) as archive:
                xml_parts = [
                    name
                    for name in archive.namelist()
                    if name.startswith("word/") and name.endswith(".xml")
                ]
                for name in xml_parts:
                    root = ElementTree.fromstring(archive.read(name))
                    text_nodes = [
                        node.text or ""
                        for node in root.iter()
                        if node.tag.endswith("}t")
                    ]
                    part_text = _clean_text(" ".join(text_nodes))
                    if part_text:
                        nonempty.append(part_text)
        except (BadZipFile, ElementTree.ParseError, KeyError) as exc:
            raise ResumeParseError(f"DOCX 内部 XML 损坏：{exc}") from exc
    return "\n".join(nonempty), max(1, len(nonempty))


def _read_pdf(path: Path) -> tuple[str, int]:
    try:
        import fitz
    except ImportError as exc:
        raise ResumeParseError("缺少 PyMuPDF，请先安装项目依赖。") from exc

    try:
        with fitz.open(path) as document:
            if document.needs_pass:
                raise ResumeParseError("PDF 已加密，无法解析。")
            if len(document) > MAX_PDF_PAGES:
                raise ResumeParseError(f"PDF 不能超过 {MAX_PDF_PAGES} 页。")
            pages = [_clean_text(page.get_text("text")) for page in document]
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"PDF 文件损坏或无法读取：{exc}") from exc
    return "\n".join(page for page in pages if page), max(1, len(pages))


def parse_resume(file_path: str | Path) -> ParsedResume:
    """Parse one local resume without sending its content to any external service."""

    path = Path(file_path).expanduser()
    if not path.exists():
        raise ResumeParseError(f"文件不存在：{path}")
    if not path.is_file():
        raise ResumeParseError(f"路径不是文件：{path}")
    if path.stat().st_size == 0:
        raise ResumeParseError("文件为空。")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ResumeParseError("文件不能超过 10 MB。")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ResumeParseError(f"不支持 {suffix or '无扩展名'}；仅支持 TXT、DOCX、PDF。")

    signature = path.read_bytes()[:8]
    if suffix == ".pdf" and not signature.startswith(b"%PDF-"):
        raise ResumeParseError("文件扩展名是 PDF，但内容不是有效 PDF。")
    if suffix == ".docx" and not signature.startswith(b"PK"):
        raise ResumeParseError("文件扩展名是 DOCX，但内容不是有效 Word 文档。")
    if suffix == ".txt" and b"\x00" in path.read_bytes()[:4096]:
        raise ResumeParseError("TXT 文件包含二进制内容，已拒绝解析。")

    readers = {".txt": _read_txt, ".docx": _read_docx, ".pdf": _read_pdf}
    text, unit_count = readers[suffix](path)
    if not text:
        raise ResumeParseError("文件中没有提取到可用文本；扫描版 PDF 后续需要 OCR。")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ResumeParseError(f"提取文本超过 {MAX_EXTRACTED_CHARACTERS:,} 字符，已拒绝处理。")

    return ParsedResume(
        source_name=path.name,
        file_type=suffix.removeprefix("."),
        text=text,
        unit_count=unit_count,
        character_count=len(text),
    )

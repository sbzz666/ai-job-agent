from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from docx import Document
import pytest

from app.services.resume_parser import ResumeParseError, parse_resume


def test_parse_utf8_txt_and_clean_whitespace(tmp_path: Path) -> None:
    source = tmp_path / "resume.txt"
    source.write_text("张三  \n\nPython   FastAPI\n", encoding="utf-8")

    result = parse_resume(source)

    assert result.file_type == "txt"
    assert result.text == "张三\nPython FastAPI"
    assert result.source_name == "resume.txt"
    assert result.character_count == len(result.text)


def test_parse_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("项目经历")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技术"
    table.cell(0, 1).text = "Python"
    document.save(source)

    result = parse_resume(source)

    assert "项目经历" in result.text
    assert "技术 | Python" in result.text
    assert result.file_type == "docx"


def test_parse_docx_text_box_xml_fallback(tmp_path: Path) -> None:
    source = tmp_path / "text-box-resume.docx"
    content_types = """<?xml version="1.0" encoding="UTF-8"?>
    <Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
      <Default Extension="xml" ContentType="application/xml"/>
      <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
    </Types>"""
    relationships = """<?xml version="1.0" encoding="UTF-8"?>
    <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rId1"
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
        Target="word/document.xml"/>
    </Relationships>"""
    document_xml = """<?xml version="1.0" encoding="UTF-8"?>
    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:v="urn:schemas-microsoft-com:vml">
      <w:body><w:p><w:r><w:pict><v:shape><v:textbox><w:txbxContent>
        <w:p><w:r><w:t>文本框项目经历</w:t></w:r></w:p>
      </w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p></w:body>
    </w:document>"""
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document_xml)

    result = parse_resume(source)

    assert "文本框项目经历" in result.text


def test_parse_pdf_with_page_count(tmp_path: Path) -> None:
    source = tmp_path / "resume.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Python FastAPI")
    document.save(source)
    document.close()

    result = parse_resume(source)

    assert result.file_type == "pdf"
    assert result.unit_count == 1
    assert "Python FastAPI" in result.text


@pytest.mark.parametrize("name", ["resume.md", "resume.exe"])
def test_reject_unsupported_file(tmp_path: Path, name: str) -> None:
    source = tmp_path / name
    source.write_text("content", encoding="utf-8")

    with pytest.raises(ResumeParseError, match="不支持"):
        parse_resume(source)


def test_reject_empty_file(tmp_path: Path) -> None:
    source = tmp_path / "empty.txt"
    source.touch()

    with pytest.raises(ResumeParseError, match="文件为空"):
        parse_resume(source)


def test_reject_fake_pdf_extension(tmp_path: Path) -> None:
    source = tmp_path / "fake.pdf"
    source.write_text("这不是 PDF", encoding="utf-8")

    with pytest.raises(ResumeParseError, match="内容不是有效 PDF"):
        parse_resume(source)


def test_reject_binary_txt(tmp_path: Path) -> None:
    source = tmp_path / "binary.txt"
    source.write_bytes(b"text\x00binary")

    with pytest.raises(ResumeParseError, match="二进制"):
        parse_resume(source)
